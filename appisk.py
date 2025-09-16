import streamlit as st
import pandas as pd
import io
import urllib.parse
from openpyxl.styles import Font
from openpyxl import Workbook
import webbrowser
import time
import uuid
import re

st.title("Смета: ГЭСН, ФСБЦ, ФССЦ, ТЦ, ФЕР с М, ФОТ")

# Добавляем выбор метода расчета
method = st.radio("Выберите метод расчета:", ["РИМ (Ресурсно-индексный)", "БИМ (Базисно-индексный)"], horizontal=True)

# Добавляем поля для ввода коэффициентов
if method == "БИМ (Базисно-индексный)":
    coefficient = st.number_input("Введите общий коэффициент:", value=1.0, min_value=0.1, step=0.1)
    tc_coefficient = st.number_input("Введите коэффициент для ТЦ:", value=1.0, min_value=0.1, step=0.1)
else:
    # Для РИМ только коэффициент для ТЦ
    tc_coefficient = st.number_input("Введите коэффициент для ТЦ:", value=1.0, min_value=0.1, step=0.1)
    coefficient = 1.0

# Словарь магазинов по категориям
stores = {
    "Строительные материалы": [
        ("Спектр Томск", "https://spektr-tomsk.ru/search/?query={query}"),
        ("Стройса", "https://www.stroyssa.tomsk.ru/search/?q={query}"),
        ("СтройПарк", "https://stroyparkdiy.ru/search?search={query}"),
        ("Сатурн", "https://tomsk.saturn.net/catalog/search/?q={query}"),
        ("СтройДвор70", "https://stroydvor70.ru/search/?query={query}"),
        ("СтройДвор Томск", "https://stroydvor.tomsk.ru/search/?q{query}"),
        ("СД70", "https://sd70.ru/search/?query={query}"),
        ("АвтоСтройЛавка", "https://avtostroylavka70.ru/search/?query={query}"),
        ("М-Плюс", "https://mplus.tomsk.ru/search/?q={query}"),
        ("СтройМастер", "https://stroimaster.tomsk.ru/search/?q={query}"),
        ("SMT Group", "https://smt-group.ru/search/?q={query}"),
        ("ВсеИнструменты", "https://www.vseinstrumenti.ru/search/?q={query}"),
        ("НСК-Строй", "https://nsk-stroy.com/search/?q={query}"),
        ("Опт-Стройка", "https://opt-stroyka.ru/search/?q={query}"),
        ("Строй-Мат54", "https://stroy-mat54.ru/search/?q={query}"),
        ("LemanaPro", "https://novosibirsk.lemanapro.ru/search/?q={query}"),
        ("Ладья", "https://www.ladya.pro/search/?q={query}"),
        ("ТомскЛес", "https://tomskles.su/search/?q={query}"),
        ("Амбар-С", "https://tomsk.ambar-s.ru/search/?q={query}")
    ],
    "Вода (сантехника)": [
        ("Waterman", "https://waterman-t.ru/search/?q={query}"),
        ("Водолей", "https://томск.водолей.рф/search?q={query}"),
        ("Сантехника", "https://tomsk.santehnica.ru/search/?q={query}"),
        ("Смесители Томск", "https://smesiteli-tomsk.ru/search/?q={query}")
    ],
    "Крепеж": [
        ("Крепеж70", "https://krepeg70.ru/search/?q={query}"),
        ("ПромКрепеж", "https://tomsk.promkrepezh.ru/search/?q={query}")
    ],
    "Утеплитель": [
        ("ТСТН", "https://tomsk.tstn.ru/search/?q={query}")
    ],
    "Металл": [
        ("Ferrum-N", "https://tomsk.ferrum-n.ru/search/?q={query}"),
        ("ANEP Metall", "https://kemerovo.anepmetall.ru/search/?q={query}"),
        ("SMT Group (металл)", "https://smt-group.ru/search/?q={query}"),
        ("СПК", "https://tomsk.spk.ru/search/?q={query}"),
        ("ЕвроМет", "https://ewromet.ru/search/?q={query}"),
        ("МеталлоБаза", "https://metallobazan.ru/search/?q={query}"),
        ("МеталлКомфорт", "https://metallkomfort.ru/search/?q{query}")
    ],
    "Кровля/фасады": [
        ("МеталлПрофиль", "https://tomsk.metallprofil.ru/search/?q={query}"),
        ("ArealMetall", "https://arealmetall.com/search/?q={query}"),
        ("DTM", "https://tomsk.dtm.ru/search/?q={query}"),
        ("КровельныйЦентр", "https://кровельныйцентр.рф/search?query={query}")
    ],
    "Электрика": [
        ("Сиб-А", "https://www.sib-a.ru/search/?q={query}"),
        ("RS24", "https://rs24.ru/search.html?q={query}"),
        ("ЭТМ", "https://www.etm.ru/catalog?searchValue={query}"),
        ("ChipDip", "https://www.chipdip.ru/search?searchtext={query}"),
        ("МиниМакс", "https://www.minimaks.ru/catalog/?q={query}")
    ],
    "Пиломатериал": [
        ("EcoLes", "https://ecoles-tomsk.ru/search/?q={query}"),
        ("ТД Лесной", "https://www.tdlesnoy.com/search/?q={query}")
    ],
    "Двери": [
        ("World of Doors", "https://worldofdoors.ru/search?query={query}"),
        ("Gallery of Doors", "https://galleryofdoors.ru/search?search={query}"),
        ("Фабрика Дверей", "https://фабрикадверей.рф/tomsk/search?query={query}")
    ],
    "Вентиляция": [
        ("НеваТом", "https://www.nevatom.ru/search/?q={query}"),
        ("TomVent", "https://tomvent.ru/search/?q={query}"),
        ("VentGrad70", "https://ventgrad70.ru/search/?q={query}")
    ],
    "Базовые магазины": [ 
        ("ЯндексМаркет", "https://market.yandex.ru/search?text={query}")
    ]
}

# Функции поиска магазинов
def search_all_stores(query):
    results = []
    query_encoded = urllib.parse.quote_plus(query)
    for category, stores_list in stores.items():
        for store_name, store_url in stores_list:
            search_url = store_url.format(query=query_encoded)
            results.append((category, store_name, store_url, search_url))
    return results

def search_selected_stores(query, selected_stores_list):
    results = []
    query_encoded = urllib.parse.quote_plus(query)
    for store_name, store_url in selected_stores_list:
        category = "Базовые магазины"
        for cat, stores_list in stores.items():
            if any(s[0] == store_name for s in stores_list):
                category = cat
                break
        search_url = store_url.format(query=query_encoded)
        results.append((category, store_name, store_url, search_url))
    return results

def improved_store_search(query, stores_list):
    results = []
    query_encoded = urllib.parse.quote_plus(query)
    query_lower = query.lower()
    for store_name, store_url in stores_list:
        if query_lower in store_name.lower():
            category = "Базовые магазины"
            for cat, cat_stores in stores.items():
                if any(s[0] == store_name for s in cat_stores):
                    category = cat
                    break
            search_url = store_url.format(query=query_encoded)
            results.append((category, store_name, store_url, search_url))
    return results

def open_link(url):
    webbrowser.open_new_tab(url)

def detect_base_and_type(obosnovanie):
    if obosnovanie.startswith("ГЭСНмр"): return "gesnmr", "ГЭСНмр"
    if obosnovanie.startswith("ГЭСНм"): return "gesnm", "ГЭСНм"
    if obosnovanie.startswith("ГЭСНп"): return "gesnp", "ГЭСНп"
    if obosnovanie.startswith("ГЭСНр"): return "gesnr", "ГЭСНр"
    if obosnovanie.startswith("ГЭСН"): return "gesn", "ГЭСН"
    if obosnovanie.startswith("ФЕРм"): return "ferm", "ФЕРм"
    if obosnovanie.startswith("ФЕРр"): return "ferr", "ФЕРр"
    if obosnovanie.startswith("ФЕРп"): return "ferp", "ФЕРп"
    if obosnovanie.startswith("ФЕР"): return "fer", "ФЕР"
    if obosnovanie.startswith("ФСБЦ"): return "fsbcm", "ФСБЦ"
    if obosnovanie.startswith("ФССЦ"): return "fsscm", "ФССЦ"
    if obosnovanie.startswith("ТЦ"): return "tc", "ТЦ"
    return "", "Неизвестно"

def get_tc_links(naimenovanie):
    if not naimenovanie: return []
    query = urllib.parse.quote_plus(naimenovanie)
    links = []
    for store_name, store_url in st.session_state.selected_stores:
        if '{query}' in store_url:
            links.append(store_url.format(query=query))
        else:
            links.append(store_url)
    return links

# Сайдбар выбор магазинов
st.sidebar.header("Выбор магазинов")

if 'selected_stores' not in st.session_state:
    st.session_state.selected_stores = []

select_all = st.sidebar.checkbox("Выбрать все магазины", value=False, key="select_all")

for category, stores_list in stores.items():
    st.sidebar.subheader(category)
    category_key = f"cat_all_{category}"
    if category_key not in st.session_state:
        st.session_state[category_key] = False
    category_select_all = st.sidebar.checkbox(f"Выбрать все в {category}", value=st.session_state[category_key], key=category_key)
    if select_all or category_select_all:
        for store in stores_list:
            if store not in st.session_state.selected_stores:
                st.session_state.selected_stores.append(store)
    else:
        for store in stores_list:
            store_key = f"{category}_{store[0]}"
            if store_key not in st.session_state:
                st.session_state[store_key] = False
            if st.sidebar.checkbox(store[0], value=st.session_state[store_key], key=store_key):
                if store not in st.session_state.selected_stores:
                    st.session_state.selected_stores.append(store)
            else:
                if store in st.session_state.selected_stores:
                    st.session_state.selected_stores.remove(store)
        st.sidebar.header("Поиск по магазинам")
search_query = st.sidebar.text_input("Введите запрос для поиска:", key="search_query")
search_type = st.sidebar.radio("Поиск в:", ["Всех магазинах", "Выбранных магазинах"], key="search_type")

if search_query:
    if search_type == "Всех магазинах":
        all_stores = [(name, url) for category, stores_list in stores.items() for name, url in stores_list]
        search_results = improved_store_search(search_query, all_stores)
        st.sidebar.subheader("Результаты поиска (все магазины):")
    else:
        search_results = improved_store_search(search_query, st.session_state.selected_stores)
        st.sidebar.subheader("Результаты поиска (выбранные магазины):")
    if search_results:
        all_selected = all(any(s[0] == store_name for s in st.session_state.selected_stores) for _, store_name, _, _ in search_results)
        select_all_found = st.sidebar.checkbox("Выбрать все найденные", key="searchbox_select_all_found", value=all_selected)
        if select_all_found and not all_selected:
            for _, store_name, store_url, _ in search_results:
                if not any(s[0] == store_name for s in st.session_state.selected_stores):
                    st.session_state.selected_stores.append((store_name, store_url))
                st.session_state[f"searchbox_{store_name}"] = True
        elif not select_all_found and all_selected:
            for _, store_name, _, _ in search_results:
                st.session_state.selected_stores = [s for s in st.session_state.selected_stores if s[0] != store_name]
                st.session_state[f"searchbox_{store_name}"] = False
        for category, store_name, store_url, search_url in search_results:
            store_key = f"searchbox_{store_name}"
            selected = st.sidebar.checkbox(f"{store_name} ({category})", key=store_key, value=any(s[0] == store_name for s in st.session_state.selected_stores))
            if selected:
                if not any(s[0] == store_name for s in st.session_state.selected_stores):
                    st.session_state.selected_stores.append((store_name, store_url))
            else:
                st.session_state.selected_stores = [s for s in st.session_state.selected_stores if s[0] != store_name]

# --- Основная часть с Excel ---
uploaded_file = st.file_uploader("Выберите файл Excel", type=["xlsx", "xls"])
if uploaded_file is not None:
    df = pd.read_excel(uploaded_file)
    results = []
    equipment_results = []  # Для хранения оборудования

    current = None
    position_number = 1
    equipment_position_number = 1
    
    # Обработка данных в зависимости от выбранного метода
    if method == "РИМ (Ресурсно-индексный)":
        # Существующая логика для РИМ
        for _, row in df.iterrows():
            obosnovanie = str(row['Обоснование']).strip() if pd.notna(row['Обоснование']) else ""
            naimenovanie = str(row['Наименование работ и затрат']).strip() if pd.notna(row['Наименование работ и затрат']) else ""
            total_current_price = row.get('всего в текущем уровне цен', 0) if pd.notna(row.get('всего в текущем уровне цен', 0)) else 0

            # Проверяем, является ли строка оборудованием (содержит "О" в номере позиции)
            poziciya = str(row.get('№ п/п', '')).strip()
            is_equipment = 'О' in poziciya and any(char.isdigit() for char in poziciya)
            
            if any([obosnovanie.startswith(x) for x in ["ГЭСН", "ФСБЦ", "ФССЦ", "ТЦ", "ФЕР"]]):
                if current is not None:
                    if current['Тип'] in ["ФСБЦ", "ФССЦ", "ТЦ"]:
                        current['Ссылки'] = get_tc_links(current['Наименование'])
                    else:
                        current['Ссылки'] = []
                    
                    # Добавляем в соответствующую категорию
                    if is_equipment:
                        equipment_results.append(current)
                    else:
                        results.append(current)
                
                base, tip = detect_base_and_type(obosnovanie)
                
                current = {
                    '№ п/п': equipment_position_number if is_equipment else position_number, 
                    'Тип': tip, 
                    'Обоснование': obosnovanie,
                    'Наименование': naimenovanie, 
                    'Ед. измерения': row.get('Единица измерения', ''),
                    'Количество': row.get('на единицу измерения', 0), 
                    'Цена за ед.': row.get('на единицу измерения', 0),
                    'Текущие цены': "", 
                    'М': 0, 
                    'ФОТ': 0, 
                    'Всего в текущем уровне цен': total_current_price,
                    'Категория': 'Оборудование' if is_equipment else 'Материалы/Работы'
                }
                
                if is_equipment:
                    equipment_position_number += 1
                else:
                    position_number += 1
                continue
                
            if current and current['Тип'].startswith("ГЭСН"):
                if naimenovanie.upper() == 'М': current['М'] = total_current_price; continue
                if naimenovanie.upper() == 'ФОТ': current['ФОT'] = total_current_price; continue
                if naimenovanie.upper() == 'ВСЕГО ПО ПОЗИЦИИ': current['Всего в текущем уровне цен'] = total_current_price; continue

        if current is not None:
            if current['Тип'] in ["ФСБЦ", "ФССЦ", "ТЦ"]: 
                current['Ссылки'] = get_tc_links(current['Наименование'])
            else: 
                current['Ссылки'] = []
            
            if current.get('Категория') == 'Оборудование':
                equipment_results.append(current)
            else:
                results.append(current)

    else:  # БИМ (Базисно-индексный метод)
        current_work = None
        fot_values = {}
        m_values = {}
        
        # Сначала собираем все значения ФОТ и М
        for _, row in df.iterrows():
            obosnovanie = str(row['Обоснование']).strip() if pd.notna(row['Обоснование']) else ""
            naimenovanie = str(row['Наименование работ и затрат']).strip() if pd.notna(row['Наименование работ и затрат']) else ""
            
            if any([obosnovanie.startswith(x) for x in ["ГЭСН", "ФЕР"]]):
                # Это работа - запоминаем обоснование
                current_work = obosnovanie
                fot_values[current_work] = 0
                m_values[current_work] = 0
            
            if current_work:
                if naimenovanie.upper() == 'ФОТ':
                    fot_value = row.get('Сметная стоимость в текущем уровне цен, руб.', 0)
                    if pd.notna(fot_value):
                        fot_values[current_work] = fot_value
                
                if naimenovanie.upper() == 'М':
                    m_value = row.get('всего(сметная стоимость)', 0)
                    if pd.notna(m_value):
                        m_values[current_work] = m_value
    
        # Теперь обрабатываем основные строки
        for _, row in df.iterrows():
            obosnovanie = str(row['Обоснование']).strip() if pd.notna(row['Обоснование']) else ""
            naimenovanie = str(row['Наименование работ и затрат']).strip() if pd.notna(row['Наименование работ и затрат']) else ""
            
            # Проверяем, является ли строка оборудованием (содержит "О" в номере позиции)
            poziciya = str(row.get('№ п/п', '')).strip()
            is_equipment = 'О' in poziciya and any(char.isdigit() for char in poziciya)
            
            # Определяем тип элемента (работа или материал)
            is_work = any([obosnovanie.startswith(x) for x in ["ГЭСН", "ФЕР"]])
            is_material = any([obosnovanie.startswith(x) for x in ["ФСБЦ", "ФССЦ", "ТЦ"]])
            
            if is_work or is_material:
                if current is not None:
                    if current['Тип'] in ["ФСБЦ", "ФССЦ", "ТЦ"]:
                        current['Ссылки'] = get_tc_links(current['Наименование'])
                    else:
                        current['Ссылки'] = []
                    
                    # Добавляем в соответствующую категорию
                    if current.get('Категория') == 'Оборудование':
                        equipment_results.append(current)
                    else:
                        results.append(current)
                
                base, tip = detect_base_and_type(obosnovanie)
                
                # Для БИМ берем количество из колонки "всего с учетом коэффициентов(количество)"
                kolichestvo = row.get('всего с учетом коэффициентов(количество)', 0) if pd.notna(row.get('всего с учетом коэффициентов(количество)', 0)) else 0
                
                # Определяем сметную стоимость в зависимости от типа
                if is_work:
                    # Для работ используем значение ФОТ из собранных данных
                    smetnaya_stoimost = fot_values.get(obosnovanie, 0)
                    item_type = "Работы"
                    m_value = m_values.get(obosnovanie, 0)
                else:
                    # Для материалов берем из колонки "всего(сметная стоимость)"
                    smetnaya_stoimost = row.get('всего(сметная стоимость)', 0) if pd.notna(row.get('всего(сметная стоимость)', 0)) else 0
                    
                    # Применяем коэффициент в зависимости от типа
                    if obosnovanie.startswith("ТЦ"):
                        smetnaya_stoimost *= tc_coefficient
                    else:
                        smetnaya_stoimost *= coefficient
                    
                    item_type = "Материалы"
                    m_value = 0
                
                current = {
                    '№ п/п': equipment_position_number if is_equipment else position_number, 
                    'Тип': tip, 
                    'Обоснование': obosnovanie,
                    'Наименование': naimenovanie, 
                    'Ед. измерения': row.get('Единица измерения', ''),
                    'Количество': kolichestvo,
                    'Сметная цена': smetnaya_stoimost,
                    'Текущие цены': "",
                    'Категория': 'Оборудование' if is_equipment else item_type,
                    'М': m_value
                }
                
                if is_equipment:
                    equipment_position_number += 1
                else:
                    position_number += 1
                continue

        if current is not None:
            if current['Тип'] in ["ФСБЦ", "ФССЦ", "ТЦ"]: 
                current['Ссылки'] = get_tc_links(current['Наименование'])
            else: 
                current['Ссылки'] = []
            
            if current.get('Категория') == 'Оборудование':
                equipment_results.append(current)
            else:
                results.append(current)

    df_results = pd.DataFrame(results)
    
    # Для РИМ вычисляем сметные цены
    if method == "РИМ (Ресурсно-индексный)":
        smetnye_ceny = []
        for _, row in df_results.iterrows():
            if str(row['Тип']).startswith("ГЭСН"): 
                smetnye_ceny.append(row['ФОТ'])
            else: 
                # Для ТЦ применяем коэффициент
                if str(row['Тип']).startswith("ТЦ"):
                    smetnye_ceny.append(row['Всего в текущем уровне цен'] * tc_coefficient)
                else:
                    smetnye_ceny.append(row['Всего в текущем уровне цен'])
        df_results['Сметная цена'] = smetnye_ceny
        df_results['Количество'] = df_results['Количество'].fillna(0)
    else:
        # Для БИМ уже есть сметная цена в данных
        df_results['Количество'] = df_results['Количество'].fillna(0)

    # Обрабатываем оборудование отдельно
    if equipment_results:
        df_equipment = pd.DataFrame(equipment_results)
        
        # Для РИМ вычисляем сметные цены для оборудования
        if method == "РИМ (Ресурсно-индексный)":
            smetnye_ceny_equip = []
            for _, row in df_equipment.iterrows():
                smetnye_ceny_equip.append(row['Всего в текущем уровне цен'])
            df_equipment['Сметная цена'] = smetnye_ceny_equip
            df_equipment['Количество'] = df_equipment['Количество'].fillna(0)
        else:
            # Для БИМ уже есть сметная цена в данных
            df_equipment['Количество'] = df_equipment['Количество'].fillna(0)

    # --- Показ результатов ---
    st.header("Результаты сметы")
    
    # Вкладки для разных категорий
    tab1, tab2 = st.tabs(["Материалы и Работы", "Оборудование"])
    
    with tab1:
        for _, row in df_results.iterrows():
            with st.expander(f"{row['№ п/п']}. {row['Наименование']} ({row['Тип']})"):
                st.write(f"**Обоснование:** {row['Обоснование']}")
                st.write(f"**Ед. измерения:** {row['Ед. измерения']}")
                st.write(f"**Количество:** {row['Количество']}")
                st.write(f"**Сметная цена:** {row['Сметная цена']}")
                if str(row['Тип']).startswith("ТЦ"):
                    st.write(f"**Применен коэффициент ТЦ:** {tc_coefficient}")
                if 'Категория' in row:
                    st.write(f"**Категория:** {row['Категория']}")
                if 'Ссылки' in row and row['Ссылки']:
                    st.write("**Поиск товаров в магазинах:**")
                    for i, (store_name, store_url) in enumerate(st.session_state.selected_stores):
                        if i < len(row['Ссылки']):
                            link = row['Ссылки'][i]
                            st.markdown(f"[{store_name}]({link})", unsafe_allow_html=True)
                else:
                    st.write("Нет ссылок для поиска товаров")
    
    with tab2:
        if equipment_results:
            for _, row in df_equipment.iterrows():
                with st.expander(f"{row['№ п/п']}. {row['Наименование']} ({row['Тип']})"):
                    st.write(f"**Обоснование:** {row['Обоснование']}")
                    st.write(f"**Ед. измерения:** {row['Ед. измерения']}")
                    st.write(f"**Количество:** {row['Количество']}")
                    st.write(f"**Сметная цена:** {row['Сметная цена']}")
                    if 'Категория' in row:
                        st.write(f"**Категория:** {row['Категория']}")
                    if 'Ссылки' in row and row['Ссылки']:
                        st.write("**Поиск товаров в магазинах:**")
                        for i, (store_name, store_url) in enumerate(st.session_state.selected_stores):
                            if i < len(row['Ссылки']):
                                link = row['Ссылки'][i]
                                st.markdown(f"[{store_name}]({link})", unsafe_allow_html=True)
                    else:
                        st.write("Нет ссылок для поиска товаров")
        else:
            st.write("Оборудование не найдено")

    # --- Создание Excel ---
    output = io.BytesIO()
    wb = Workbook()
    wb.remove(wb.active)
    
    if method == "РИМ (Ресурсно-индексный)":
        # Существующая логика для РИМ
        df_gesn = df_results[df_results['Тип'].str.startswith("ГЭСН")].copy()
        df_gesn = df_gesn[['№ п/п','Тип','Обоснование','Наименование','Ед. измерения','Количество','Сметная цена','Текущие цены']]
        ws_gesn = wb.create_sheet('Работы')
        for col_idx, column_name in enumerate(df_gesn.columns, 1):
            ws_gesn.cell(row=1, column=col_idx, value=column_name)
        for row_idx, row_data in enumerate(df_gesn.itertuples(index=False), 2):
            for col_idx, value in enumerate(row_data, 1):
                ws_gesn.cell(row=row_idx, column=col_idx, value=value)

        df_other = df_results[~df_results['Тип'].str.startswith("ГЭСН")].copy()
        store_names = [store[0] for store in st.session_state.selected_stores]
        for i, (store_name, store_url) in enumerate(st.session_state.selected_stores):
            df_other[store_name] = df_other['Ссылки'].apply(lambda x: x[i] if isinstance(x, list) and len(x)>i else "")
        df_other = df_other.drop(columns=['Ссылки'])
        df_other = df_other[['№ п/п','Тип','Обоснование','Наименование','Ед. измерения','Количество','Сметная цена','Текущие цены'] + store_names]

        ws_other = wb.create_sheet('Материалы')
        for col_idx, column_name in enumerate(df_other.columns,1):
            ws_other.cell(row=1,column=col_idx,value=column_name)
        for row_idx, row_data in enumerate(df_other.itertuples(index=False),2):
            for col_idx, value in enumerate(row_data,1):
                column_name = df_other.columns[col_idx-1]
                if column_name in store_names and value and isinstance(value,str) and value.startswith("http"):
                    cell = ws_other.cell(row=row_idx, column=col_idx, value=column_name)
                    cell.hyperlink = value
                    cell.font = Font(color="0000FF", underline="single")
                else:
                    ws_other.cell(row=row_idx, column=col_idx, value=value)

        gesn_sum_m = df_results[df_results['Тип'].str.startswith("ГЭСН")]['М'].sum()
        last_row = ws_other.max_row + 2
        ws_other.cell(row=last_row, column=6).value = "Итого М"
        ws_other.cell(row=last_row, column=7).value = gesn_sum_m

    else:
        # Новая логика для БИМ - разделение на Работы и Материалы
        df_works = df_results[df_results['Категория'] == 'Работы'].copy()
        df_works = df_works[['№ п/п','Тип','Обоснование','Наименование','Ед. измерения','Количество','Сметная цена','Текущие цены']]
        ws_works = wb.create_sheet('Работы')
        for col_idx, column_name in enumerate(df_works.columns, 1):
            ws_works.cell(row=1, column=col_idx, value=column_name)
        for row_idx, row_data in enumerate(df_works.itertuples(index=False), 2):
            for col_idx, value in enumerate(row_data, 1):
                ws_works.cell(row=row_idx, column=col_idx, value=value)

        df_materials = df_results[df_results['Категория'] == 'Материалы'].copy()
        store_names = [store[0] for store in st.session_state.selected_stores]
        for i, (store_name, store_url) in enumerate(st.session_state.selected_stores):
            df_materials[store_name] = df_materials['Ссылки'].apply(lambda x: x[i] if isinstance(x, list) and len(x)>i else "")
        df_materials = df_materials.drop(columns=['Ссылки', 'Категория'])
        df_materials = df_materials[['№ п/п','Тип','Обоснование','Наименование','Ед. измерения','Количество','Сметная цена','Текущие цены'] + store_names]

        ws_materials = wb.create_sheet('Материалы')
        for col_idx, column_name in enumerate(df_materials.columns,1):
            ws_materials.cell(row=1,column=col_idx,value=column_name)
        for row_idx, row_data in enumerate(df_materials.itertuples(index=False),2):
            for col_idx, value in enumerate(row_data,1):
                column_name = df_materials.columns[col_idx-1]
                if column_name in store_names and value and isinstance(value,str) and value.startswith("http"):
                    cell = ws_materials.cell(row=row_idx, column=col_idx, value=column_name)
                    cell.hyperlink = value
                    cell.font = Font(color="0000FF", underline="single")
                else:
                    ws_materials.cell(row=row_idx, column=col_idx, value=value)
        
        # Добавляем итого М как в РИМ - суммируем М для всех работ
        bim_sum_m = df_results[df_results['Категория'] == 'Работы']['М'].sum()
        last_row = ws_materials.max_row + 2
        ws_materials.cell(row=last_row, column=6).value = "Итого М"
        ws_materials.cell(row=last_row, column=7).value = bim_sum_m

    # Добавляем лист для оборудования
    if equipment_results:
        df_equipment_export = df_equipment.copy()
        store_names = [store[0] for store in st.session_state.selected_stores]
        for i, (store_name, store_url) in enumerate(st.session_state.selected_stores):
            df_equipment_export[store_name] = df_equipment_export['Ссылки'].apply(lambda x: x[i] if isinstance(x, list) and len(x)>i else "")
        df_equipment_export = df_equipment_export.drop(columns=['Ссылки', 'Категория'])
        df_equipment_export = df_equipment_export[['№ п/п','Тип','Обоснование','Наименование','Ед. измерения','Количество','Сметная цена','Текущие цены'] + store_names]

        ws_equipment = wb.create_sheet('Оборудование')
        for col_idx, column_name in enumerate(df_equipment_export.columns,1):
            ws_equipment.cell(row=1,column=col_idx,value=column_name)
        for row_idx, row_data in enumerate(df_equipment_export.itertuples(index=False),2):
            for col_idx, value in enumerate(row_data,1):
                column_name = df_equipment_export.columns[col_idx-1]
                if column_name in store_names and value and isinstance(value,str) and value.startswith("http"):
                    cell = ws_equipment.cell(row=row_idx, column=col_idx, value=column_name)
                    cell.hyperlink = value
                    cell.font = Font(color="0000FF", underline="single")
                else:
                    ws_equipment.cell(row=row_idx, column=col_idx, value=value)

    wb.save(output)
    output.seek(0)

    st.success(f"Выбрано магазинов: {len(st.session_state.selected_stores)}")
    
    # Ввод имени файла пользователем
    filename_input = st.text_input("Введите имя файла для сохранения:", value="")

    # Если пользователь ничего не ввёл, можно дать подсказку
    if not filename_input.strip():
        st.warning("Введите имя файла для скачивания")
    else:
        # Добавляем расширение .xlsx если его нет
        if not filename_input.lower().endswith('.xlsx'):
            filename = filename_input + '.xlsx'
        else:
            filename = filename_input

        st.download_button(
            label="Скачать Excel",
            data=output.getvalue(),
            file_name=filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"download_{filename}"
        )

# --- Выбранные магазины в сайдбаре ---
st.sidebar.header("Выбранные магазины")
if st.session_state.selected_stores:
    for store in st.session_state.selected_stores:
        st.sidebar.write(f"✓ {store[0]}")
else:
    st.sidebar.write("Магазины не выбраны")